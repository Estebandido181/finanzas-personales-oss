"""Lector del Excel de "Descargar movimientos" de Bancolombia (ítem 28, D31).

Convierte los bytes del archivo .xlsx que el dueño descarga de la Sucursal
Virtual en una lista de ``MovimientoExcel`` (fecha, monto con signo en
centavos enteros, descripción y destino). Es una función pura: no guarda el
archivo y no toca la base de datos ni la red. De la columna ``Referencia``
(celulares o cuentas de terceros) conserva solo los últimos 4 dígitos, como
``*1234`` (D37); el número completo nunca sale de este módulo: no se devuelve,
no aparece en mensajes de error ni en registros.

Si el archivo no coincide EXACTAMENTE con ``ESTRUCTURA`` no se adivina: se
lanza ``ExcelNoReconocido`` con un motivo en español que nunca incluye datos
del archivo (ni valores de celdas ni el nombre del archivo), solo la fila o la
columna donde falló.

Seguridad (el archivo viene de afuera y no es confiable), en este orden:

1. Tamaño máximo del archivo ANTES de abrirlo (``TAMANO_MAXIMO_BYTES``).
2. Solo .xlsx: debe ser un ZIP (firma ``PK``) cuyo ``[Content_Types].xml``
   declare un libro .xlsx normal. Se rechazan .xlsm (con macros), plantillas,
   .xls antiguos, CSV y cualquier otra cosa, y también las partes de macros
   (``vbaProject``), controles ActiveX, hojas de macros y vínculos externos
   aunque vengan con la extensión correcta. Si se pasa ``nombre_archivo``, su
   extensión también debe ser ``.xlsx``.
3. Bombas zip: se limita la cantidad de partes del ZIP y el tamaño
   descomprimido REAL (se cuenta al descomprimir, no se confía en lo que
   declara el ZIP). Las partes cifradas se rechazan.
4. Partes: solo se aceptan partes .xml y .rels, más imágenes .png, .jpg,
   .jpeg o .gif con la firma real de su formato (no se pueden analizar como
   XML). Cualquier otra extensión se rechaza, porque openpyxl decide qué
   parsear por las relaciones del libro y no por la extensión: así toda parte
   que pueda llegar al analizador XML pasa por las revisiones 5 y 6.
5. XML: se rechaza cualquier parte con ``<!DOCTYPE`` o ``<!ENTITY`` (ataques
   de expansión de entidades y entidades externas; un .xlsx normal nunca los
   trae) o con bytes nulos (XML en UTF-16/32, que esa revisión no vería).
6. Fórmulas: se RECHAZA el archivo completo si alguna parte trae un elemento
   ``<f>`` (fórmula). Con ``data_only=True`` openpyxl no evalúa fórmulas, pero
   devolvería el último valor que guardó quien editó el archivo, y ese valor
   no se puede verificar. Un archivo descargado del banco no trae fórmulas: si
   las tiene, alguien lo editó, y tratar solo esa celda como inválida dejaría
   pasar el resto de un archivo alterado. Por eso se rechaza entero.
7. openpyxl se abre con ``read_only=True``, ``data_only=True``,
   ``keep_vba=False`` y ``keep_links=False``; openpyxl nunca ejecuta macros.
   Se ignoran las dimensiones que declara la hoja (pueden mentir) y se cuentan
   las filas al leerlas: más de ``FILAS_MAXIMAS`` -> rechazo.

Límite conocido: openpyxl usa el analizador XML de la biblioteca estándar
(expat >= 2.4 ya limita la expansión de entidades) y no ``defusedxml``, que
no se agrega para no sumar dependencias; las revisiones de los puntos 4 y 5 cubren esos
ataques antes de que openpyxl lea el archivo.
"""

from __future__ import annotations

import io
import math
import re
import unicodedata
import warnings
import zipfile
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from finanzas.modulos.importaciones.contratos import MovimientoExcel, ExcelNoReconocido

# ---------------------------------------------------------------------------
# Estructura esperada del archivo
# ---------------------------------------------------------------------------
# Estructura confirmada (ítem 28) con un Excel inventado que el dueño armó con
# la misma forma que el del banco (hoy lo genera tests/excel_inventado.py),
# sin subir el real:
#   - una sola hoja, llamada exactamente "Hoja 1";
#   - encabezados en la fila 1: Fecha, Descripción, Referencia, Valor; sin
#     filas de título ni de totales; el banco corta cada descarga en un tope
#     de filas que no se conoce (se han visto 146, D51);
#   - Fecha: celda de fecha de Excel con hora 00:00 o 05:00 (ver _fecha);
#   - Descripción: texto recortado a unos 30 caracteres, a veces con dos
#     espacios seguidos ("COMPRA EN  TIENDA ..."), que se conservan;
#   - Referencia: vacía o texto de dígitos (celulares o cuentas); se conservan
#     solo los últimos 4 dígitos (D37). Si trae otra cosa (letras, menos de 4
#     dígitos), la fila entra sin destino y se marca (D45); no se rechaza;
#   - Valor: celda numérica (entero o decimal) con signo; sin saldo.
# Si algo no coincide, el lector rechaza el archivo y dice qué.
ESTRUCTURA = {
    "hoja": "Hoja 1",  # nombre exacto; debe ser la única hoja.
    "fila_encabezados": 1,  # los movimientos empiezan en la fila siguiente.
    "encabezados": ("Fecha", "Descripción", "Referencia", "Valor"),
    # Fecha: celda de fecha de Excel; solo se usa el día. Hora permitida:
    # 00:00 o 05:00 (HORAS_FECHA_PERMITIDAS); cualquier otra se rechaza.
    "fecha": "celda de fecha de Excel con hora 00:00 o 05:00",
    # Valor: celda numérica con signo (negativo = sale de la cuenta, positivo =
    # entra) y como máximo 2 decimales. Nunca texto, nunca cero.
    "valor": "celda numérica con signo y hasta 2 decimales",
}

# El banco no da la hora del movimiento, solo el día. El archivo trae las
# fechas con hora 05:00: es la medianoche de Colombia (UTC-5) escrita en UTC,
# así que el día es el mismo y basta con tomar la fecha. Se acepta también
# 00:00 (la fecha sola). Cualquier otra hora significaría otra conversión de
# zona horaria que no conocemos, en la que el día podría correrse: en vez de
# adivinar, se rechaza el archivo.
HORAS_FECHA_PERMITIDAS = (time(0, 0), time(5, 0))

_HOJA: str = ESTRUCTURA["hoja"]  # type: ignore[assignment]
_FILA_ENCABEZADOS: int = ESTRUCTURA["fila_encabezados"]  # type: ignore[assignment]
_ENCABEZADOS: tuple[str, ...] = ESTRUCTURA["encabezados"]  # type: ignore[assignment]
_COL_FECHA, _COL_DESCRIPCION, _COL_REFERENCIA, _COL_VALOR = range(4)

# ---------------------------------------------------------------------------
# Límites de seguridad y de cordura
# ---------------------------------------------------------------------------
# Un archivo real de un par de cientos de movimientos pesa unas decenas de KB; los límites dejan
# margen amplio sin permitir archivos que agoten memoria o CPU.
TAMANO_MAXIMO_BYTES = 2 * 1024 * 1024  # archivo .xlsx tal como se sube
DESCOMPRIMIDO_MAXIMO_BYTES = 10 * 1024 * 1024  # suma real de todas las partes
PARTES_MAXIMAS = 100  # partes (archivos) dentro del ZIP
FILAS_MAXIMAS = 1000  # filas de la hoja, encabezado y filas vacías incluidos
LONGITUD_MAXIMA_DESCRIPCION = 100  # el banco recorta a unos 30 caracteres
# Valor absoluto máximo: $10.000.000.000,00. Más que eso es un dato corrupto.
MONTO_MAXIMO_CENTAVOS = 10_000_000_000 * 100
ANIO_MINIMO = 2000
ANIO_MAXIMO = 2099

_FIRMA_ZIP = b"PK\x03\x04"
_TIPO_LIBRO_XLSX = (
    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
)
# Partes que un .xlsx descargado del banco no trae y que pueden ejecutar o
# traer contenido de afuera. Se comparan en minúsculas contra la ruta.
_PARTES_PROHIBIDAS = ("vbaproject", "activex", "macrosheets", "externallinks")
# Elemento de fórmula de SpreadsheetML, con o sin prefijo de espacio de
# nombres: <f>, <f t="shared" ...>, <f/>, <x:f>. No coincide con <formula> ni
# <filterColumn> porque exige que después de "f" venga espacio, "/" o ">".
_ELEMENTO_FORMULA = re.compile(rb"<(?:[A-Za-z_][\w.\-]*:)?f[\s/>]")
_DECLARACIONES_XML = (b"<!DOCTYPE", b"<!ENTITY")
# Únicas partes no XML permitidas: imágenes (por ejemplo, un logo), con la
# firma real de su formato.
_FIRMAS_IMAGEN = {
    "png": b"\x89PNG\r\n\x1a\n",
    "jpg": b"\xff\xd8\xff",
    "jpeg": b"\xff\xd8\xff",
    "gif": b"GIF8",
}






# ---------------------------------------------------------------------------
# Revisión del contenedor antes de abrirlo con openpyxl
# ---------------------------------------------------------------------------


def _revisar_parte(ruta: str, datos: bytes) -> None:
    """Revisa una parte ya descomprimida según su extensión.

    Solo se aceptan partes .xml y .rels (a las que se aplican las tres
    revisiones de XML) e imágenes .png, .jpg, .jpeg o .gif cuyos primeros bytes
    son de verdad los de esa imagen. Una imagen así nunca se puede analizar
    como XML (no empieza por "<"), aunque alguien la declare como hoja; y en
    modo ``read_only`` openpyxl no lee imágenes. Cualquier otra extensión
    (.dat, .bin, sin extensión...) se rechaza: openpyxl elige qué parsear
    según las relaciones, no según la extensión, y así nada queda sin revisar.
    """
    extension = ruta.lower().rsplit("/", 1)[-1].rpartition(".")[2]
    if ruta.endswith("/") and not datos:
        return  # entrada de carpeta vacía del ZIP
    if extension in _FIRMAS_IMAGEN:
        if not datos.startswith(_FIRMAS_IMAGEN[extension]):
            raise ExcelNoReconocido(
                "El archivo trae una imagen que no corresponde a su formato."
            )
        return
    if extension not in ("xml", "rels"):
        raise ExcelNoReconocido(
            "El archivo trae partes internas de un tipo no permitido. Solo se "
            "aceptan .xlsx descargados del banco, sin editar."
        )
    if b"\x00" in datos:
        raise ExcelNoReconocido("El archivo usa una codificación interna no soportada.")
    if any(d in datos for d in _DECLARACIONES_XML):
        raise ExcelNoReconocido("El archivo trae declaraciones XML no permitidas.")
    if _ELEMENTO_FORMULA.search(datos):
        raise ExcelNoReconocido(
            "El archivo tiene fórmulas. Un reporte descargado del "
            "banco no las trae; descárgalo de nuevo sin editarlo."
        )


def _revisar_contenedor(contenido: bytes) -> None:
    """Revisa el ZIP sin interpretar el libro. Lanza ``ExcelNoReconocido``."""
    if not contenido.startswith(_FIRMA_ZIP):
        raise ExcelNoReconocido(
            "El archivo no es un Excel .xlsx. Descarga los movimientos en "
            "formato Excel (.xlsx); no se aceptan .xls, .xlsm, CSV ni PDF."
        )
    try:
        archivo_zip = zipfile.ZipFile(io.BytesIO(contenido))
    except (zipfile.BadZipFile, ValueError, OSError, EOFError):
        raise ExcelNoReconocido("El archivo .xlsx está dañado o incompleto.") from None

    with archivo_zip:
        partes = archivo_zip.infolist()
        if len(partes) > PARTES_MAXIMAS:
            raise ExcelNoReconocido(
                "El archivo tiene demasiadas partes internas para ser un "
                "reporte de movimientos."
            )
        if sum(p.file_size for p in partes) > DESCOMPRIMIDO_MAXIMO_BYTES:
            raise ExcelNoReconocido(
                "El archivo descomprimido es demasiado grande para ser un "
                "reporte de movimientos."
            )

        tipos_contenido = None
        restante = DESCOMPRIMIDO_MAXIMO_BYTES
        for parte in partes:
            ruta = parte.filename
            if parte.flag_bits & 0x1:
                raise ExcelNoReconocido(
                    "El archivo está cifrado o protegido con contraseña."
                )
            if any(p in ruta.lower() for p in _PARTES_PROHIBIDAS):
                raise ExcelNoReconocido(
                    "El archivo trae macros, controles o vínculos externos. "
                    "Solo se aceptan .xlsx descargados del banco, sin editar."
                )
            try:
                with archivo_zip.open(parte) as origen:
                    # Se lee como máximo lo que queda del presupuesto + 1 byte:
                    # si llega a sobrar, el tamaño declarado mentía.
                    datos = origen.read(restante + 1)
            except (zipfile.BadZipFile, NotImplementedError, ValueError, OSError, EOFError):
                raise ExcelNoReconocido(
                    "El archivo .xlsx está dañado o usa una compresión no soportada."
                ) from None
            if len(datos) > restante:
                raise ExcelNoReconocido(
                    "El archivo descomprimido es demasiado grande para ser un "
                    "reporte de movimientos."
                )
            restante -= len(datos)

            _revisar_parte(ruta, datos)
            if ruta == "[Content_Types].xml":
                tipos_contenido = datos

    if tipos_contenido is None or _TIPO_LIBRO_XLSX not in tipos_contenido:
        # Un .xlsm declara "macroEnabled.main+xml"; una plantilla,
        # "template.main+xml". Ninguno trae el tipo de un .xlsx normal.
        raise ExcelNoReconocido(
            "El archivo no es un libro .xlsx normal (por ejemplo, es un .xlsm "
            "con macros o una plantilla)."
        )


# ---------------------------------------------------------------------------
# Conversión de celdas
# ---------------------------------------------------------------------------


def _vacia(valor: object) -> bool:
    return valor is None or (isinstance(valor, str) and valor.strip() == "")


def _texto_encabezado(valor: object) -> str | None:
    if not isinstance(valor, str):
        return None
    # NFC para que "Descripción" escrita con tilde combinada también coincida;
    # no cambia letras ni mayúsculas.
    return unicodedata.normalize("NFC", valor.strip())


def _fecha(valor: object, fila: int) -> date:
    # openpyxl entrega datetime para celdas con formato de fecha. Un número sin
    # formato de fecha, un texto o una hora sola no se aceptan.
    if not isinstance(valor, datetime):
        raise ExcelNoReconocido(
            f"Fila {fila}: la columna Fecha no es una fecha de Excel."
        )
    if valor.time() not in HORAS_FECHA_PERMITIDAS or valor.tzinfo is not None:
        raise ExcelNoReconocido(
            f"Fila {fila}: la columna Fecha trae una hora inesperada; se "
            "esperaba 00:00 o 05:00."
        )
    if not ANIO_MINIMO <= valor.year <= ANIO_MAXIMO:
        raise ExcelNoReconocido(
            f"Fila {fila}: la columna Fecha está fuera del rango de años "
            f"{ANIO_MINIMO}-{ANIO_MAXIMO}."
        )
    return valor.date()


def monto_a_centavos(valor: object) -> int | None:
    """Convierte el número de una celda a centavos enteros con signo, sin floats.

    - ``int`` -> exacto.
    - ``float`` -> ``Decimal(repr(valor))``: ``repr`` da la representación más
      corta que vuelve al mismo float (``-5000.5`` -> ``"-5000.5"``), que es
      el número escrito en el archivo. Nunca se multiplica un float.
    - Más de 2 decimales, cero, infinito, NaN, booleanos, textos o un valor
      absoluto mayor que ``MONTO_MAXIMO_CENTAVOS`` -> ``None`` (no se
      redondea ni se adivina).
    """
    if isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        numero = Decimal(valor)
    elif isinstance(valor, float):
        if not math.isfinite(valor):
            return None
        try:
            numero = Decimal(repr(valor))
        except InvalidOperation:
            return None
    else:
        return None
    if abs(numero) > Decimal(MONTO_MAXIMO_CENTAVOS) / 100:
        return None
    centavos = numero.scaleb(2)
    if centavos != centavos.to_integral_value():
        return None
    resultado = int(centavos)
    if resultado == 0:
        return None
    return resultado


_REFERENCIA = re.compile(r"[0-9]{4,}")
# Más que esto no es un celular ni una cuenta; se trata como sin formato.
LONGITUD_MAXIMA_REFERENCIA = 40


def ultimos_4_de_referencia(valor: object) -> tuple[str | None, bool]:
    """``(destino, sin_formato)`` de una celda de ``Referencia`` (D37, D45).

    - vacía -> ``(None, False)``;
    - texto (o entero) de solo dígitos ASCII, al menos 4 -> ``("*" + últimos
      4, False)``. Los espacios al inicio y al final se ignoran;
    - cualquier otra cosa (letras, signos, espacios internos, menos de 4
      dígitos, decimales, fechas) -> ``(None, True)``: la fila entra sin
      destino. No se adivina.

    El valor completo nunca se devuelve ni se guarda en ninguna parte.
    """
    if _vacia(valor):
        return None, False
    if isinstance(valor, bool):
        return None, True
    if isinstance(valor, int):
        texto = str(valor) if valor >= 0 else ""
    elif isinstance(valor, str):
        texto = valor.strip()
    else:
        return None, True
    if len(texto) > LONGITUD_MAXIMA_REFERENCIA or not _REFERENCIA.fullmatch(texto):
        return None, True
    return "*" + texto[-4:], False


def _descripcion(valor: object, fila: int) -> str:
    if not isinstance(valor, str) or valor.strip() == "":
        raise ExcelNoReconocido(
            f"Fila {fila}: la columna Descripción está vacía o no es texto."
        )
    texto = valor.strip()
    if len(texto) > LONGITUD_MAXIMA_DESCRIPCION:
        raise ExcelNoReconocido(
            f"Fila {fila}: la columna Descripción es demasiado larga."
        )
    if any(unicodedata.category(c) == "Cc" for c in texto):
        raise ExcelNoReconocido(
            f"Fila {fila}: la columna Descripción trae caracteres de control."
        )
    return texto


# ---------------------------------------------------------------------------
# Lectura del libro
# ---------------------------------------------------------------------------


def _filas(contenido: bytes) -> list[tuple[int, tuple, bool]]:
    """Abre el libro ya revisado y devuelve sus filas ya leídas.

    Cada fila es ``(número, celdas de las columnas esperadas, hay_sobrantes)``.
    Todo se lee dentro del bloque que silencia los warnings y el libro se
    cierra antes de devolver: nada queda abierto mientras se validan las filas.
    """
    # Importación diferida: el resto de la app no carga openpyxl si no lo usa.
    import openpyxl

    # openpyxl avisa con warnings que pueden incluir valores de celdas; se
    # silencian para que nada del archivo llegue a los registros.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            libro = openpyxl.load_workbook(
                io.BytesIO(contenido),
                read_only=True,
                data_only=True,
                keep_vba=False,
                keep_links=False,
            )
        except Exception:
            raise ExcelNoReconocido(
                "El archivo no se pudo leer como un libro de Excel .xlsx."
            ) from None
        try:
            if libro.sheetnames != [_HOJA]:
                raise ExcelNoReconocido(
                    f'El archivo debe tener una sola hoja, llamada "{_HOJA}".'
                )
            hoja = libro[_HOJA]
            if not hasattr(hoja, "reset_dimensions"):
                raise ExcelNoReconocido(
                    f'La hoja "{_HOJA}" no es una hoja de datos.'
                )
            # Las dimensiones declaradas pueden mentir (por ejemplo A1:XFD1048576)
            # y openpyxl rellenaría millones de celdas vacías: se ignoran.
            hoja.reset_dimensions()
            columnas = len(_ENCABEZADOS)
            filas: list[tuple[int, tuple, bool]] = []
            try:
                for numero, fila in enumerate(
                    hoja.iter_rows(min_row=1, values_only=True), start=1
                ):
                    if numero > FILAS_MAXIMAS:
                        raise ExcelNoReconocido(
                            f"El archivo tiene más de {FILAS_MAXIMAS} filas; "
                            "es demasiado grande para ser un reporte de movimientos."
                        )
                    fila = tuple(fila)
                    # Solo se guardan las columnas esperadas y si sobró algo:
                    # una fila puede venir con miles de celdas vacías de relleno.
                    celdas = fila[:columnas] + (None,) * max(0, columnas - len(fila))
                    hay_sobrantes = any(not _vacia(v) for v in fila[columnas:])
                    filas.append((numero, celdas, hay_sobrantes))
            except ExcelNoReconocido:
                raise
            except Exception:
                raise ExcelNoReconocido(
                    "El archivo no se pudo leer como un libro de Excel .xlsx."
                ) from None
            return filas
        finally:
            libro.close()


def leer_excel_movimientos(
    contenido: bytes, nombre_archivo: str | None = None
) -> list[MovimientoExcel]:
    """Lee el Excel de "Descargar movimientos" y devuelve sus movimientos.

    ``contenido``: bytes del archivo tal como se subió. ``nombre_archivo``
    (opcional): si se da, su extensión debe ser ``.xlsx``; nunca aparece en
    los mensajes de error.

    Devuelve los movimientos en el orden del archivo (lista vacía si solo trae
    encabezados). Lanza ``ExcelNoReconocido`` si el archivo no es seguro de
    leer o no coincide exactamente con ``ESTRUCTURA``.
    """
    if not isinstance(contenido, (bytes, bytearray)):
        raise ExcelNoReconocido("No se recibió el contenido del archivo.")
    if len(contenido) > TAMANO_MAXIMO_BYTES:
        raise ExcelNoReconocido(
            "El archivo es demasiado grande para ser un reporte de movimientos "
            f"(máximo {TAMANO_MAXIMO_BYTES // (1024 * 1024)} MB)."
        )
    if nombre_archivo is not None and not nombre_archivo.lower().endswith(".xlsx"):
        raise ExcelNoReconocido(
            "Solo se aceptan archivos .xlsx (no .xlsm, .xls, CSV ni otros)."
        )
    contenido = bytes(contenido)
    _revisar_contenedor(contenido)

    columnas = len(_ENCABEZADOS)
    movimientos: list[MovimientoExcel] = []
    vio_encabezados = False
    fila_vacia_previa: int | None = None

    for numero, celdas, hay_sobrantes in _filas(contenido):
        if numero < _FILA_ENCABEZADOS:
            if hay_sobrantes or not all(_vacia(v) for v in celdas):
                raise ExcelNoReconocido(
                    f"Se esperaban los encabezados en la fila {_FILA_ENCABEZADOS}."
                )
            continue

        if hay_sobrantes:
            raise ExcelNoReconocido(
                f"Fila {numero}: hay datos fuera de las {columnas} columnas esperadas."
            )

        if numero == _FILA_ENCABEZADOS:
            if tuple(_texto_encabezado(v) for v in celdas) != _ENCABEZADOS:
                raise ExcelNoReconocido(
                    f"Los encabezados de la fila {_FILA_ENCABEZADOS} no coinciden; "
                    f"se esperaban, en orden: {', '.join(_ENCABEZADOS)}."
                )
            vio_encabezados = True
            continue

        if all(_vacia(v) for v in celdas):
            # Solo se aceptan filas vacías al final: una fila vacía entre
            # movimientos indica otra estructura (por ejemplo, totales abajo).
            if fila_vacia_previa is None:
                fila_vacia_previa = numero
            continue
        if fila_vacia_previa is not None:
            raise ExcelNoReconocido(
                f"Fila {fila_vacia_previa}: hay una fila vacía entre movimientos."
            )

        fecha = _fecha(celdas[_COL_FECHA], numero)
        descripcion = _descripcion(celdas[_COL_DESCRIPCION], numero)
        # De Referencia solo salen los últimos 4 dígitos (D37); nunca aparece
        # en mensajes. Sin formato no rechaza el archivo (D45).
        referencia, sin_formato = ultimos_4_de_referencia(celdas[_COL_REFERENCIA])
        centavos = monto_a_centavos(celdas[_COL_VALOR])
        if centavos is None:
            raise ExcelNoReconocido(
                f"Fila {numero}: la columna Valor no es un número distinto de "
                "cero con máximo 2 decimales."
            )
        movimientos.append(
            MovimientoExcel(fecha, centavos, descripcion, referencia, sin_formato)
        )

    if not vio_encabezados:
        raise ExcelNoReconocido(
            f"Se esperaban los encabezados en la fila {_FILA_ENCABEZADOS}."
        )
    return movimientos
