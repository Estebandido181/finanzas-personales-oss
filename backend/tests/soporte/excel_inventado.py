"""Excel INVENTADO de "Descargar movimientos" generado por código (ítem 40, D43).

Desde D43 no se suben archivos ``.xlsx`` al repositorio: ``revisar_datos.py``
no mira dentro de ellos y habría que revisarlos a mano. Este módulo arma en
memoria, al correr las pruebas, el mismo Excel inventado que el dueño preparó
para el ítem 28 con la forma exacta del reporte del banco:

* una sola hoja llamada "Hoja 1", encabezados en la fila 1;
* ``Fecha`` como celda de fecha con hora 05:00 (la medianoche de Colombia
  guardada en UTC) y formato ``mm-dd-yy``;
* ``Descripción`` recortada, a veces con dos espacios seguidos;
* ``Referencia`` vacía o como texto de dígitos (celulares o cuentas, todos
  inventados);
* ``Valor`` numérico con signo y centavos; sin saldo.

Totales que las pruebas comprueban: 63 movimientos que suman -33.180.181
centavos (gastos -74.183.000 y abonos 41.002.819).

Uso::

    from tests.soporte.excel_inventado import generar_excel, MOVIMIENTOS_INVENTADOS
    contenido = generar_excel()                      # los 63 movimientos
    contenido = generar_excel(filas=[...])           # otras filas inventadas

Para la conciliación (ítems 29 a 31) se pueden pasar otras filas con la misma
forma ``(fecha, descripcion, referencia, valor)``; ``fecha`` puede ser ``date``
o ``datetime`` (un ``date`` se guarda a las 05:00, como lo hace el banco).
"""

from __future__ import annotations

import io
from datetime import date, datetime, time

from openpyxl import Workbook

ENCABEZADOS = ("Fecha", "Descripción", "Referencia", "Valor")
HOJA = "Hoja 1"
FORMATO_FECHA = "mm-dd-yy"
HORA_DEL_BANCO = time(5, 0)  # medianoche de Colombia (UTC-5) guardada en UTC

TOTAL_CENTAVOS = -33180181
GASTOS_CENTAVOS = -74183000
ABONOS_CENTAVOS = 41002819

# (mes, día, descripción, referencia, valor) del año 2026, en el orden del
# banco (más reciente primero). Todo inventado: comercios "EJEMPLO",
# celulares 30000000NN y cuentas en ceros.
_FILAS = [
    (8, 30, 'ABONO INTERESES AHORROS', None, 1.35),
    (8, 29, 'ABONO INTERESES AHORROS', None, 1.08),
    (8, 28, 'ABONO INTERESES AHORROS', None, 1.06),
    (8, 28, 'COBRO VIAJE TRANSPORTE MASIVO', None, -7100),
    (8, 27, 'ABONO INTERESES AHORROS', None, 0.92),
    (8, 27, 'COBRO VIAJE TRANSPORTE MASIVO', None, -10650),
    (8, 27, 'TRANSFERENCIAS A NEQUI', '3000000004', -6200),
    (8, 26, 'ABONO INTERESES AHORROS', None, 1.14),
    (8, 26, 'COBRO VIAJE TRANSPORTE MASIVO', None, -10650),
    (8, 25, 'ABONO INTERESES AHORROS', None, 1.13),
    (8, 25, 'COBRO VIAJE TRANSPORTE MASIVO', None, -3550),
    (8, 25, 'COMPRA EN  RESTAURANTE EJEMPLO', None, -31500),
    (8, 24, 'ABONO INTERESES AHORROS', None, 0.91),
    (8, 24, 'COBRO VIAJE TRANSPORTE MASIVO', None, -3550),
    (8, 23, 'ABONO INTERESES AHORROS', None, 1.14),
    (8, 22, 'ABONO INTERESES AHORROS', None, 0.86),
    (8, 22, 'TRANSFERENCIA CTA SUC VIRTUAL', '00000000002', 160000),
    (8, 21, 'ABONO INTERESES AHORROS', None, 0.99),
    (8, 21, 'COBRO VIAJE TRANSPORTE MASIVO', None, -3550),
    (8, 20, 'ABONO INTERESES AHORROS', None, 1.12),
    (8, 20, 'COBRO VIAJE TRANSPORTE MASIVO', None, -10650),
    (8, 20, 'TRANSF A PERSONA EJEMPLO DOS', None, -27000),
    (8, 19, 'ABONO INTERESES AHORROS', None, 0.97),
    (8, 19, 'COBRO VIAJE TRANSPORTE MASIVO', None, -3550),
    (8, 18, 'ABONO INTERESES AHORROS', None, 0.83),
    (8, 18, 'COBRO VIAJE TRANSPORTE MASIVO', None, -3550),
    (8, 18, 'COMPRA EN  WOMPI EJEMPLO', None, -9800),
    (8, 17, 'ABONO INTERESES AHORROS', None, 1.04),
    (8, 17, 'COBRO VIAJE TRANSPORTE MASIVO', None, -3550),
    (8, 16, 'ABONO INTERESES AHORROS', None, 1.15),
    (8, 15, 'ABONO INTERESES AHORROS', None, 1.37),
    (8, 15, 'COMPRA EN  MERCADO EJEMPLO', None, -68630),
    (8, 15, 'TRANSFERENCIAS A NEQUI', '3000000003', -3500),
    (8, 14, 'ABONO INTERESES AHORROS', None, 1.18),
    (8, 14, 'COBRO VIAJE TRANSPORTE MASIVO', None, -10650),
    (8, 13, 'ABONO INTERESES AHORROS', None, 0.87),
    (8, 13, 'COBRO VIAJE TRANSPORTE MASIVO', None, -3550),
    (8, 12, 'ABONO INTERESES AHORROS', None, 0.84),
    (8, 12, 'COBRO VIAJE TRANSPORTE MASIVO', None, -10650),
    (8, 12, 'TRANSFERENCIA CTA SUC VIRTUAL', '00000000001', 250000),
    (8, 11, 'ABONO INTERESES AHORROS', None, 0.85),
    (8, 11, 'COBRO VIAJE TRANSPORTE MASIVO', None, -7100),
    (8, 10, 'ABONO INTERESES AHORROS', None, 1.05),
    (8, 10, 'COBRO VIAJE TRANSPORTE MASIVO', None, -3550),
    (8, 10, 'PAGO QR COMERCIO EJEMPLO', None, -12000),
    (8, 9, 'ABONO INTERESES AHORROS', None, 0.85),
    (8, 8, 'ABONO INTERESES AHORROS', None, 0.93),
    (8, 8, 'RETIRO CAJERO CAJERO EJEMPLO', None, -50000),
    (8, 7, 'ABONO INTERESES AHORROS', None, 0.83),
    (8, 7, 'COBRO VIAJE TRANSPORTE MASIVO', None, -10650),
    (8, 6, 'ABONO INTERESES AHORROS', None, 0.86),
    (8, 6, 'COBRO VIAJE TRANSPORTE MASIVO', None, -10650),
    (8, 6, 'PAGO PSE EMPRESA EJEMPLO SAS', None, -350000),
    (8, 5, 'ABONO INTERESES AHORROS', None, 0.84),
    (8, 5, 'COBRO VIAJE TRANSPORTE MASIVO', None, -10650),
    (8, 5, 'TRANSF A PERSONA EJEMPLO UNO', None, -10000),
    (8, 4, 'ABONO INTERESES AHORROS', None, 1.04),
    (8, 4, 'COBRO VIAJE TRANSPORTE MASIVO', None, -3550),
    (8, 4, 'TRANSFERENCIAS A NEQUI', '3000000001', -6200),
    (8, 4, 'TRANSFERENCIAS A NEQUI', '3000000002', -6200),
    (8, 3, 'ABONO INTERESES AHORROS', None, 0.99),
    (8, 3, 'COBRO VIAJE TRANSPORTE MASIVO', None, -3550),
    (8, 3, 'COMPRA EN  TIENDA EJEMPLO', None, -25900),]

MOVIMIENTOS_INVENTADOS: list[tuple[date, str, str | None, int | float]] = [
    (date(2026, mes, dia), descripcion, referencia, valor)
    for mes, dia, descripcion, referencia, valor in _FILAS
]


def generar_excel(filas=None, *, hoja: str = HOJA, encabezados=ENCABEZADOS) -> bytes:
    """Arma el .xlsx inventado en memoria y devuelve sus bytes.

    ``filas``: lista de ``(fecha, descripcion, referencia, valor)``; por defecto
    los 63 ``MOVIMIENTOS_INVENTADOS``. Nada se escribe en disco.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    ws.append(list(encabezados))
    for fecha, descripcion, referencia, valor in (
        MOVIMIENTOS_INVENTADOS if filas is None else filas
    ):
        if isinstance(fecha, date) and not isinstance(fecha, datetime):
            fecha = datetime.combine(fecha, HORA_DEL_BANCO)
        ws.append([fecha, descripcion, referencia, valor])
        ws.cell(row=ws.max_row, column=1).number_format = FORMATO_FECHA
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
